from __future__ import annotations

import json
import tempfile
from pathlib import Path
from typing import Iterable

import pandas as pd
from docx import Document as DocxDocument
from pypdf import PdfReader

from app.config import SUPPORTED_EXTENSIONS
from services.models import LoadedDocument
from services.utils import make_doc_id, normalize_text


class UnsupportedFileError(ValueError):
    pass


def load_uploaded_file(file) -> LoadedDocument:
    suffix = Path(file.name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedFileError(f"Unsupported file type '{suffix}'. Supported: {supported}")

    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(file.getvalue())
        tmp_path = Path(tmp.name)

    try:
        text, metadata = load_path(tmp_path, display_name=file.name)
    finally:
        tmp_path.unlink(missing_ok=True)

    return LoadedDocument(
        doc_id=make_doc_id(file.name, text),
        name=file.name,
        source_type=suffix.lstrip("."),
        text=normalize_text(text),
        metadata=metadata,
    )


def load_text_input(text: str, name: str = "Pasted Text") -> LoadedDocument:
    cleaned = normalize_text(text)
    return LoadedDocument(
        doc_id=make_doc_id(name, cleaned),
        name=name,
        source_type="text",
        text=cleaned,
        metadata={"characters": len(cleaned)},
    )


def load_path(path: Path, display_name: str | None = None) -> tuple[str, dict]:
    suffix = path.suffix.lower()
    name = display_name or path.name
    if suffix == ".pdf":
        return _load_pdf(path), {"filename": name}
    if suffix == ".docx":
        return _load_docx(path), {"filename": name}
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore"), {"filename": name}
    if suffix == ".csv":
        df = pd.read_csv(path)
        return _dataframe_to_text(df, "CSV"), {"filename": name, "tables": [_table_metadata(df, "CSV")]}
    if suffix in {".xlsx", ".xls"}:
        return _load_excel(path, name)
    if suffix == ".json":
        text, tables = _load_json(path)
        metadata = {"filename": name}
        if tables:
            metadata["tables"] = tables
        return text, metadata
    raise UnsupportedFileError(f"Unsupported file type: {suffix}")


def _load_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages: list[str] = []
    for idx, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(f"[Page {idx}]\n{text}")
    return "\n\n".join(pages)


def _load_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    parts: list[str] = []
    parts.extend(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())
    for table_idx, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip() for cell in row.cells))
        if rows:
            parts.append(f"[Table {table_idx}]\n" + "\n".join(rows))
    return "\n\n".join(parts)


def _load_excel(path: Path, filename: str) -> tuple[str, dict]:
    sheets = pd.read_excel(path, sheet_name=None)
    parts = []
    tables = []
    for sheet_name, df in sheets.items():
        parts.append(_dataframe_to_text(df, f"Sheet: {sheet_name}"))
        tables.append(_table_metadata(df, f"Sheet: {sheet_name}"))
    return "\n\n".join(parts), {"filename": filename, "tables": tables}


def _load_json(path: Path) -> tuple[str, list[dict]]:
    with path.open("r", encoding="utf-8", errors="ignore") as handle:
        data = json.load(handle)
    tables = []
    if isinstance(data, list) and all(isinstance(item, dict) for item in data):
        tables.append(_table_metadata(pd.DataFrame(data), "JSON records"))
    elif isinstance(data, dict):
        for key, value in data.items():
            if isinstance(value, list) and all(isinstance(item, dict) for item in value):
                tables.append(_table_metadata(pd.DataFrame(value), f"JSON {key}"))
    return json.dumps(data, ensure_ascii=False, indent=2), tables


def _dataframe_to_text(df: pd.DataFrame, label: str) -> str:
    preview = df.fillna("").astype(str)
    records: Iterable[str] = (
        "; ".join(f"{col}: {row[col]}" for col in preview.columns) for _, row in preview.iterrows()
    )
    return f"[{label}]\nColumns: {', '.join(map(str, preview.columns))}\n" + "\n".join(records)


def _table_metadata(df: pd.DataFrame, label: str) -> dict:
    clean = df.fillna("")
    return {
        "label": label,
        "columns": [str(column) for column in clean.columns],
        "rows": clean.astype(object).to_dict(orient="records"),
        "row_count": int(len(clean)),
    }
