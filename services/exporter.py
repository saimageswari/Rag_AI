from __future__ import annotations

import warnings
from datetime import datetime
from pathlib import Path

from docx import Document as DocxDocument

warnings.filterwarnings("ignore", message="You have both PyFPDF & fpdf2 installed.*")
from fpdf import FPDF

from app.config import EXPORT_DIR
from services.models import ChatTurn, ExportFilter


FILTER_LABELS = {
    "file_only": "File Q&A only",
    "ai_only": "Explanation only",
    "full": "Full conversation",
}


def export_conversation(turns: list[ChatTurn], export_filter: ExportFilter, fmt: str) -> Path:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = EXPORT_DIR / f"smart_doc_conversation_{export_filter}_{timestamp}.{fmt}"
    entries = _render_entries(turns, export_filter)
    if fmt == "txt":
        path.write_text(_as_text(entries, export_filter), encoding="utf-8")
    elif fmt == "docx":
        _write_docx(path, entries, export_filter)
    elif fmt == "pdf":
        _write_pdf(path, entries, export_filter)
    else:
        raise ValueError(f"Unsupported export format: {fmt}")
    return path


def _render_entries(turns: list[ChatTurn], export_filter: ExportFilter) -> list[dict]:
    rendered = []
    for turn in turns:
        entry = {
            "timestamp": turn.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
            "question": turn.question,
            "file_answer": turn.file_answer if export_filter in {"file_only", "full"} else None,
            "ai_answer": turn.ai_answer if export_filter in {"ai_only", "full"} else None,
            "sources": turn.sources if export_filter in {"file_only", "full"} else [],
        }
        if export_filter == "ai_only" and not turn.ai_answer:
            continue
        rendered.append(entry)
    return rendered


def _as_text(entries: list[dict], export_filter: ExportFilter) -> str:
    lines = [f"AI Smart Document Assistant Export - {FILTER_LABELS[export_filter]}", ""]
    for idx, entry in enumerate(entries, start=1):
        lines.extend([f"Q{idx}. {entry['question']}", f"Timestamp: {entry['timestamp']}"])
        if entry.get("file_answer"):
            lines.extend(["", "File-grounded answer:", entry["file_answer"]])
        if entry.get("ai_answer"):
            lines.extend(["", "Explanation:", entry["ai_answer"]])
        if entry.get("sources"):
            source_labels = [
                f"{source.source} chunk {source.chunk_id}" + (f" page {source.page}" if source.page else "")
                for source in entry["sources"]
            ]
            lines.extend(["", "Sources:", "; ".join(source_labels)])
        lines.extend(["", "-" * 72, ""])
    return "\n".join(lines)


def _write_docx(path: Path, entries: list[dict], export_filter: ExportFilter) -> None:
    doc = DocxDocument()
    doc.add_heading("AI Smart Document Assistant Export", 0)
    doc.add_paragraph(FILTER_LABELS[export_filter])
    for idx, entry in enumerate(entries, start=1):
        doc.add_heading(f"Q{idx}. {entry['question']}", level=1)
        doc.add_paragraph(f"Timestamp: {entry['timestamp']}")
        if entry.get("file_answer"):
            doc.add_heading("File-grounded answer", level=2)
            doc.add_paragraph(entry["file_answer"])
        if entry.get("ai_answer"):
            doc.add_heading("Explanation", level=2)
            doc.add_paragraph(entry["ai_answer"])
        if entry.get("sources"):
            doc.add_heading("Sources", level=2)
            for source in entry["sources"]:
                page = f", page {source.page}" if source.page else ""
                doc.add_paragraph(f"{source.source}{page}, chunk {source.chunk_id}", style="List Bullet")
    doc.save(path)


def _write_pdf(path: Path, entries: list[dict], export_filter: ExportFilter) -> None:
    pdf = FPDF()
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    width = _pdf_width(pdf)
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(width, 9, "Smart Chat Assistant Export")
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(width, 7, FILTER_LABELS[export_filter])
    pdf.ln(4)
    for idx, entry in enumerate(entries, start=1):
        pdf.set_font("Helvetica", "B", 12)
        pdf.multi_cell(width, 7, _pdf_safe(f"Q{idx}. {entry['question']}"))
        pdf.set_font("Helvetica", "", 9)
        pdf.multi_cell(width, 6, f"Timestamp: {entry['timestamp']}")
        if entry.get("file_answer"):
            _pdf_section(pdf, "File-grounded answer", entry["file_answer"])
        if entry.get("ai_answer"):
            _pdf_section(pdf, "Explanation", entry["ai_answer"])
        if entry.get("sources"):
            sources = "; ".join(
                f"{source.source} chunk {source.chunk_id}" + (f" page {source.page}" if source.page else "")
                for source in entry["sources"]
            )
            _pdf_section(pdf, "Sources", sources)
        pdf.ln(4)
    pdf.output(str(path))


def _pdf_section(pdf: FPDF, title: str, body: str) -> None:
    width = _pdf_width(pdf)
    pdf.set_font("Helvetica", "B", 10)
    pdf.multi_cell(width, 6, title)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(width, 6, _pdf_safe(body))


def _pdf_safe(text: str) -> str:
    return text.encode("latin-1", errors="replace").decode("latin-1")


def _pdf_width(pdf: FPDF) -> float:
    return max(120, pdf.w - pdf.l_margin - pdf.r_margin)
